"""
Verificación del MVP: recalcula los números DESDE EL XLSX de forma independiente
(sin pasar por el ETL) y los compara contra lo que responde la API.

Sirve para el criterio de aceptación de R1-01: «los KPI cuadran contra una suma
directa en Excel».

Uso:
    python -m backend.tests.verificar                  # base local + xlsx
    python -m backend.tests.verificar http://127.0.0.1:8000   # también la API
"""

from __future__ import annotations

import sys
import urllib.request
import json
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from backend.app import config, db, queries  # noqa: E402
from backend.app.queries import Filtros  # noqa: E402

RAIZ = Path(__file__).resolve().parents[2]
fallos: list[str] = []
pruebas = 0


def check(nombre: str, obtenido, esperado, tol: float = 0) -> None:
    global pruebas
    pruebas += 1
    ok = abs(obtenido - esperado) <= tol if isinstance(esperado, (int, float)) else obtenido == esperado
    print(f"  {'✓' if ok else '✗'} {nombre}: {obtenido!r}" + ("" if ok else f"  (esperado {esperado!r})"))
    if not ok:
        fallos.append(f"{nombre}: obtenido {obtenido!r}, esperado {esperado!r}")


# --------------------------------------------------------------------------- #
# 1. Recálculo independiente desde el xlsx
# --------------------------------------------------------------------------- #
def desde_xlsx(ruta: Path) -> dict:
    """Suma directa sobre el archivo, aplicando solo la regla de filas de
    continuación (una sección puede ocupar varias filas y la fila extra no trae
    columna Total)."""
    wb = load_workbook(ruta, read_only=True, data_only=True)
    res = {"secciones": 0, "asignados": 0, "cupo": 0, "por_hoja": {}, "cursos": set(),
           "docentes": set(), "sobrecupo": 0, "cero": 0}
    for hoja in ("Clase", "Laboratorio"):
        ws, n = wb[hoja], 0
        for r in ws.iter_rows(min_row=8, values_only=True):
            if not r[1] or r[25] is None:
                continue
            n += 1
            total, limite = int(r[25] or 0), int(r[26] or 0)
            res["asignados"] += total
            res["cupo"] += limite
            res["cursos"].add(str(r[1]).strip().zfill(4))
            if r[17]:
                res["docentes"].add(" ".join(str(r[17]).split()))
            if limite and total > limite:
                res["sobrecupo"] += 1
            if total == 0:
                res["cero"] += 1
        res["por_hoja"][hoja] = n
        res["secciones"] += n
    wb.close()
    return res


# --------------------------------------------------------------------------- #
# 2. Comparación contra la réplica local
# --------------------------------------------------------------------------- #
def main() -> int:
    archivos = sorted(db.SNAPSHOTS_DIR.glob("*.xlsx"))
    if not archivos:
        print(f"No hay snapshots en {db.SNAPSHOTS_DIR}")
        return 1
    ruta = archivos[0]

    print(f"\n[1] Recálculo independiente sobre {ruta.name}")
    x = desde_xlsx(ruta)
    print(f"    xlsx → {x['secciones']} secciones {x['por_hoja']}, "
          f"{x['asignados']:,} asignados, cupo {x['cupo']:,}")

    con = db.connect()
    snap = queries.snapshot_vigente(con)
    if not snap:
        print("La base está vacía. Ejecute: python -m backend.app.etl")
        return 1
    f = Filtros(snapshot_id=snap["id"])

    print("\n[2] Réplica local (SQLite) vs xlsx")
    k = queries.kpis(con, f)
    check("total de secciones", k["secciones"], x["secciones"])
    check("total de asignados", k["asignados"], x["asignados"])
    check("cupo total", k["cupo"], x["cupo"])
    check("cursos distintos", k["cursos"], len(x["cursos"]))
    check("catedráticos distintos", k["docentes"], len(x["docentes"]))

    for hoja, n in x["por_hoja"].items():
        fh = Filtros(snapshot_id=snap["id"], tipo=hoja)
        check(f"secciones de {hoja}", queries.kpis(con, fh)["secciones"], n)

    print("\n[3] Coherencia interna")
    o = queries.ocupacion(con, f)
    check("las categorías de estado suman el total de secciones",
          sum(e["secciones"] for e in o["por_estado"]), k["secciones"])
    check("el histograma suma el total de secciones",
          sum(h["secciones"] for h in o["histograma"]), k["secciones"])
    check("secciones en sobrecupo", next(e["secciones"] for e in o["por_estado"]
                                         if e["estado"] == "Sobrecupo"), x["sobrecupo"])
    check("secciones sin asignados", next(e["secciones"] for e in o["por_estado"]
                                          if e["estado"] == "Sin asignados"), x["cero"])

    for columna in ("modalidad", "jornada", "area", "tipo"):
        d = queries.distribucion(con, f, columna)
        check(f"la distribución por {columna} suma el total de asignados",
              sum(r["asignados"] for r in d), k["asignados"])

    print("\n[4] Reglas de limpieza")
    dup = con.execute(
        "SELECT COUNT(*) FROM (SELECT seccion_id, nombre FROM miembros "
        "GROUP BY seccion_id, nombre HAVING COUNT(*) > 1)").fetchone()[0]
    check("sin auxiliares duplicados dentro de una misma sección", dup, 0)
    check("ninguna sección quedó sin bloques de horario salvo las de modalidad externa",
          con.execute("""SELECT COUNT(*) FROM secciones s
                         WHERE s.snapshot_id = ? AND s.modalidad != 'Externo / sin aula'
                           AND NOT EXISTS (SELECT 1 FROM bloques b WHERE b.seccion_id = s.id)""",
                      (snap["id"],)).fetchone()[0], 0)
    check("la Ñ corrupta del reporte quedó corregida",
          con.execute("SELECT COUNT(*) FROM secciones WHERE curso LIKE '%\\%' ESCAPE '\\'"
                      ).fetchone()[0], 0)
    check("toda sección tiene área asignada",
          con.execute("SELECT COUNT(*) FROM secciones WHERE area IS NULL").fetchone()[0], 0)
    sin_clasificar = con.execute(
        "SELECT COUNT(*) FROM secciones WHERE area = ?", (config.AREA_POR_DEFECTO,)).fetchone()[0]
    print(f"  · secciones sin clasificar por el mapa de áreas: {sin_clasificar} "
          f"({sin_clasificar / k['secciones']:.1%}) — revisar backend/app/config.py si crece")

    print("\n[5] Filtros")
    foco = Filtros(snapshot_id=snap["id"], area=config.AREA_FOCO)
    kf = queries.kpis(con, foco)
    print(f"  · {config.AREA_FOCO}: {kf['secciones']} secciones, {kf['asignados']:,} asignados, "
          f"ocupación {kf['ocupacion']:.1%}")
    check("el filtro por área es un subconjunto del total", kf["asignados"] <= k["asignados"], True)
    busq = queries.secciones(con, Filtros(snapshot_id=snap["id"], q="REDES"), limite=500)
    check("la búsqueda libre devuelve resultados", busq["total"] > 0, True)

    # --------------------------------------------------------------------- #
    if len(sys.argv) > 1:
        base = sys.argv[1].rstrip("/")
        print(f"\n[6] API en {base}")
        try:
            opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
            api = json.loads(opener.open(f"{base}/api/kpis", timeout=10).read())
            check("la API reporta el mismo total de asignados", api["asignados"], k["asignados"])
            check("la API reporta las mismas secciones", api["secciones"], k["secciones"])
            csv_txt = opener.open(f"{base}/api/export.csv", timeout=30).read().decode("utf-8")
            filas_csv = len(csv_txt.strip().splitlines()) - 6  # 5 de cabecera + encabezado
            check("el CSV exporta una fila por sección", filas_csv, k["secciones"])
        except Exception as exc:  # noqa: BLE001
            fallos.append(f"API inaccesible: {exc}")
            print(f"  ✗ API inaccesible: {exc}")

    print("\n" + "─" * 64)
    if fallos:
        print(f"FALLARON {len(fallos)} de {pruebas} verificaciones:")
        for f_ in fallos:
            print(f"  · {f_}")
        return 1
    print(f"Las {pruebas} verificaciones pasaron.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

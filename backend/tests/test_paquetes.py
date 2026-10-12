"""Paquetes .zip con varios reportes: subida web, carpeta de entrada y consola."""

from __future__ import annotations

import os
import time
import zipfile
from datetime import datetime

from openpyxl import load_workbook

from backend.app import config, etl, paquetes, vigilante
from backend.app.autoprueba import SECCIONES_CLASE, SECCIONES_LAB, construir_reporte

from .conftest import REAL_2026, subir


def escalar(secciones, factor):
    return [dict(s, regulares=round(s["regulares"] * factor), otra_fac=0) for s in secciones]


def _reportes(tmp_path, anios=(2024, 2025, 2026)):
    out = []
    for i, anio in enumerate(anios):
        out.append(construir_reporte(
            tmp_path / f"{anio}_S1.xlsx", anio=anio, corte=datetime(anio, 3, 1),
            clase=escalar(SECCIONES_CLASE, 1 + 0.1 * i), laboratorio=escalar(SECCIONES_LAB, 1 + 0.1 * i)))
    return out


def _zip(ruta, miembros: dict):
    with zipfile.ZipFile(ruta, "w", zipfile.ZIP_DEFLATED) as z:
        for nombre, contenido in miembros.items():
            z.writestr(nombre, contenido)
    return ruta


def test_xlsx_suelto_no_es_paquete():
    assert not paquetes.es_paquete(REAL_2026)


def test_subir_zip_carga_cada_reporte_y_habilita_predicciones(cliente, tmp_path):
    reps = _reportes(tmp_path)
    z = _zip(tmp_path / "semestre_2024_a_2026.zip", {
        **{f"semestre_2024_a_2026/{r.name}": r.read_bytes() for r in reps},
        "__MACOSX/semestre_2024_a_2026/._2024_S1.xlsx": b"basura de mac",
        "semestre_2024_a_2026/.DS_Store": b"x",
        "semestre_2024_a_2026/LEEME.txt": b"hola",
        "semestre_2024_a_2026/otro.zip": b"PK\x05\x06" + b"\0" * 18,
    })
    r = subir(cliente, z).json()
    por = {x["archivo"].split(" › ")[-1]: x for x in r["resultados"]}
    assert set(por) == {"2024_S1.xlsx", "2025_S1.xlsx", "2026_S1.xlsx", "LEEME.txt", "otro.zip"}
    assert all(por[f"{a}_S1.xlsx"]["resultado"] == "aceptado" for a in (2024, 2025, 2026))
    assert por["LEEME.txt"]["resultado"] == "rechazado"
    assert "otro .zip" in por["otro.zip"]["mensaje"]
    assert r["resumen"]["aceptado"] == 3
    # el nombre visible conserva de qué paquete vino
    assert por["2024_S1.xlsx"]["archivo"] == "semestre_2024_a_2026.zip › 2024_S1.xlsx"
    p = cliente.get("/api/predicciones", params={"periodo": "PRIMER SEMESTRE"}).json()
    assert p["disponible"] and p["anio_objetivo"] == 2027
    # subirlo de nuevo: todo duplicado, nada se carga dos veces
    r2 = subir(cliente, z).json()
    assert r2["resumen"]["duplicado"] == 3 and r2["resumen"]["aceptado"] == 0


def test_zip_sin_reportes(cliente, tmp_path):
    z = _zip(tmp_path / "nada.zip", {"a.txt": b"x"})
    r = subir(cliente, z).json()["resultados"]
    assert len(r) == 1 and r[0]["resultado"] == "rechazado" and "no contiene reportes" in r[0]["mensaje"]


def test_zip_que_es_un_respaldo_se_explica(cliente, tmp_path):
    z = _zip(tmp_path / "respaldo_20261012.zip", {"asignados.db": b"x", "archivo/ab/ab.xlsx": b"x"})
    r = subir(cliente, z).json()["resultados"]
    assert len(r) == 1 and r[0]["resultado"] == "rechazado" and "restaurar" in r[0]["mensaje"]


def test_zip_danado(cliente, tmp_path):
    reps = _reportes(tmp_path, (2025,))
    z = _zip(tmp_path / "p.zip", {r.name: r.read_bytes() for r in reps})
    datos = z.read_bytes()
    (tmp_path / "roto.zip").write_bytes(datos[: len(datos) // 2])
    r = subir(cliente, tmp_path / "roto.zip").json()["resultados"]
    assert len(r) == 1 and r[0]["resultado"] == "rechazado"


def test_reporte_recortado_dentro_del_zip_se_rechaza(cliente, tmp_path):
    bueno, = _reportes(tmp_path, (2025,))
    malo = construir_reporte(tmp_path / "2024_S1.xlsx", anio=2024, corte=datetime(2024, 3, 1), sin_pie=True)
    wb = load_workbook(malo)
    wb["Clase"].cell(row=wb["Clase"].max_row + 1, column=1, value="... (truncated at 1000 rows)")
    wb.save(malo)
    z = _zip(tmp_path / "p.zip", {bueno.name: bueno.read_bytes(), malo.name: malo.read_bytes()})
    por = {x["archivo"].split(" › ")[-1]: x for x in subir(cliente, z).json()["resultados"]}
    assert por["2025_S1.xlsx"]["resultado"] == "aceptado"
    assert por["2024_S1.xlsx"]["resultado"] == "rechazado"
    assert "recortada" in por["2024_S1.xlsx"]["mensaje"]


def test_miembro_demasiado_grande(cliente, tmp_path, monkeypatch):
    reps = _reportes(tmp_path, (2025,))
    z = _zip(tmp_path / "p.zip", {reps[0].name: reps[0].read_bytes()})
    monkeypatch.setattr(config, "MAX_BYTES", 1000)
    monkeypatch.setattr(config, "MAX_MB", 0.001)
    r = subir(cliente, z).json()["resultados"]
    assert r[0]["resultado"] == "rechazado" and "máximo" in r[0]["mensaje"]


def _envejecer(carpeta):
    viejo = time.time() - 300
    for raiz, _, archivos in os.walk(carpeta):
        for f in archivos:
            os.utime(os.path.join(raiz, f), (viejo, viejo))


def test_zip_en_carpeta_de_entrada(entorno, tmp_path):
    reps = _reportes(tmp_path, (2025, 2026))
    _zip(config.ENTRADA_DIR / "historia.zip", {f"h/{r.name}": r.read_bytes() for r in reps})
    _envejecer(config.ENTRADA_DIR)
    vigilante.revisar_una_vez()                       # primera vista: tamaño estable
    res = vigilante.revisar_una_vez()
    assert sorted(r.resultado for r in res) == ["aceptado", "aceptado"]
    assert not (config.ENTRADA_DIR / "historia.zip").exists()
    assert any(p.name.endswith("historia.zip") for p in config.PROCESADOS_DIR.iterdir())


def test_zip_descomprimido_dentro_de_entrada(entorno, tmp_path):
    """Si alguien descomprime el .zip dentro de entrada/, queda una subcarpeta: también se procesa."""
    sub = config.ENTRADA_DIR / "semestre_2021_a_2026"
    sub.mkdir(parents=True)
    for r in _reportes(tmp_path, (2025, 2026)):
        (sub / r.name).write_bytes(r.read_bytes())
    _envejecer(config.ENTRADA_DIR)
    res = vigilante.revisar_una_vez(exigir_estable=False)
    assert sorted(r.resultado for r in res) == ["aceptado", "aceptado"]
    assert not sub.exists()                           # la subcarpeta vacía se retira
    # las carpetas del sistema nunca se procesan
    config.RECHAZADOS_DIR.mkdir(parents=True, exist_ok=True)
    (config.RECHAZADOS_DIR / "x.xlsx").write_bytes(REAL_2026.read_bytes())
    assert vigilante.revisar_una_vez(exigir_estable=False) == []


def test_zip_con_un_rechazado_deja_motivo(entorno, tmp_path):
    bueno, = _reportes(tmp_path, (2025,))
    _zip(config.ENTRADA_DIR / "mixto.zip", {bueno.name: bueno.read_bytes(), "nota.txt": b"x"})
    res = vigilante.revisar_una_vez(exigir_estable=False)
    assert sorted(r.resultado for r in res) == ["aceptado", "rechazado"]
    motivo = next(config.PROCESADOS_DIR.glob("*mixto.zip.motivo.txt"))
    assert "nota.txt" in motivo.read_text(encoding="utf-8")


def test_consola_carga_zip(entorno, tmp_path, monkeypatch, capsys):
    monkeypatch.setattr("backend.app.privilegios.ceder", lambda: None)
    reps = _reportes(tmp_path, (2025, 2026))
    z = _zip(tmp_path / "h.zip", {r.name: r.read_bytes() for r in reps})
    assert etl.main(["cargar", str(z)]) == 0
    salida = capsys.readouterr().out
    assert salida.count("ACEPTADO") == 2


def test_archivos_con_fecha_en_el_futuro_se_procesan(entorno, tmp_path):
    """Caso real: al descomprimir un .zip, los archivos conservan la hora con la que se
    comprimieron (otra zona horaria) y quedan «en el futuro». Antes nunca se procesaban."""
    sub = config.ENTRADA_DIR / "semestre"
    sub.mkdir(parents=True)
    for r in _reportes(tmp_path, (2025, 2026)):
        destino = sub / r.name
        destino.write_bytes(r.read_bytes())
        futuro = time.time() + 6 * 3600
        os.utime(destino, (futuro, futuro))
    assert vigilante.revisar_una_vez() == []          # primera vista: se anota el tamaño
    res = vigilante.revisar_una_vez()                 # segunda: tamaño estable -> se procesa
    assert sorted(r.resultado for r in res) == ["aceptado", "aceptado"]


def test_archivo_que_no_se_puede_mover_no_se_reprocesa(entorno, tmp_path, monkeypatch):
    """Si la subcarpeta no permite mover (permisos del equipo), se carga UNA vez y no se
    vuelve a procesar en cada revisión."""
    sub = config.ENTRADA_DIR / "semestre"
    sub.mkdir(parents=True)
    r, = _reportes(tmp_path, (2025,))
    (sub / r.name).write_bytes(r.read_bytes())
    original = type(sub).replace

    def sin_permiso(self, destino):
        if self.parent == sub:
            raise PermissionError(13, "Permission denied")
        return original(self, destino)
    monkeypatch.setattr(type(sub), "replace", sin_permiso)
    primera = vigilante.revisar_una_vez(exigir_estable=False)
    assert [x.resultado for x in primera] == ["aceptado"]
    assert vigilante.revisar_una_vez(exigir_estable=False) == []
    assert vigilante.revisar_una_vez(exigir_estable=False) == []
    assert any(p.name.startswith("_AVISO_no_se_pudo_mover") for p in config.RECHAZADOS_DIR.iterdir())

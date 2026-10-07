"""Lectura y validación del archivo: formatos reales y todo lo que puede venir mal."""

from __future__ import annotations

import io
import zipfile
from datetime import datetime, time

import pytest
from openpyxl import load_workbook

from backend.app import config, lector
from backend.app.autoprueba import SECCIONES_CLASE, construir_reporte

from .conftest import REAL_2025, REAL_2026, REAL_VACIO


def codigos(rep, nivel=None):
    return {m.codigo for m in rep.mensajes if nivel is None or m.nivel == nivel}


# ═══════════════════════════ archivos reales ═══════════════════════════
def test_real_primer_semestre_2026():
    rep = lector.leer_reporte(REAL_2026)
    assert rep.es_valido and not rep.vacio
    assert rep.periodo.codigo == "01" and rep.anio == 2026
    assert rep.fecha_corte == datetime(2026, 9, 9, 13, 6)
    assert len(rep.secciones) == 1378            # 926 clase + 452 laboratorio
    assert sum(s["tipo"] == "Clase" for s in rep.secciones) == 926
    assert rep.total_asignados == 40869          # = 30,865 + 10,004 del pie
    assert rep.control["Clase"] == {"total_reg": 30865, "total": 30865}
    assert not rep.trae_clasificacion
    assert not rep.advertencias


def test_real_vacaciones_2025_formato_33_columnas():
    rep = lector.leer_reporte(REAL_2025)
    assert rep.es_valido
    assert rep.periodo.codigo == "02" and rep.anio == 2025
    assert rep.trae_clasificacion
    assert rep.total_asignados == 6921 + 2263
    assert rep.control["Laboratorio"]["total_reg"] == 2222
    cerradas = [s for s in rep.secciones if s["cerrada"]]
    assert len(cerradas) == 154
    assert all(s["estado"] == "Cerrada" and s["docente"] is None for s in cerradas)
    assert any(m["tipo"] == "RESPONSABLE CARGA DE NOTAS" for s in rep.secciones for m in s["miembros"])
    assert {s["escuela"] for s in rep.secciones} >= {"CIENCIAS Y SISTEMAS", "CIENCIAS", "CIVIL"}
    # el límite mínimo trae dato y se usa
    assert any(s["limite_min"] == 65 for s in rep.secciones)


def test_real_reporte_vacio():
    rep = lector.leer_reporte(REAL_VACIO)
    assert rep.es_valido and rep.vacio
    assert rep.periodo.codigo == "14" and rep.anio == 2023
    assert set(rep.hojas_vacias) == {"Clase", "Laboratorio"}


def test_filas_de_continuacion_no_inflan_totales():
    """En el reporte real hay 44 secciones de clase con más de un bloque de horario."""
    rep = lector.leer_reporte(REAL_2026)
    multi = [s for s in rep.secciones if len({(b["hora_inicio"], b["hora_fin"]) for b in s["bloques"]}) > 1]
    assert len(multi) >= 40
    wb = load_workbook(REAL_2026, read_only=True, data_only=True)
    filas = sum(1 for ws in wb.worksheets for r in ws.iter_rows(min_row=8, values_only=True) if r[1])
    wb.close()
    assert filas > len(rep.secciones)


def test_corrige_ene_corrupta():
    rep = lector.leer_reporte(REAL_2026)
    cursos = {s["curso"] for s in rep.secciones}
    assert not any("\\" in c for c in cursos)
    assert "ANALISIS Y DISEÑO DE SISTEMAS 1" in cursos


# ═══════════════════════════ variantes sintéticas ═══════════════════════════
def test_sintetico_valido_y_continuacion(tmp_path):
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx"))
    assert rep.es_valido, rep.mensajes
    ed = next(s for s in rep.secciones if s["codigo"] == "0772" and s["tipo"] == "Clase")
    assert {b["dia"] for b in ed["bloques"]} == {"Ma", "J", "S"}
    staff = next(s for s in rep.secciones if s["codigo"] == "0781")
    assert staff["sin_docente"] and not staff["cerrada"]


def test_columna_nueva_se_ignora_con_advertencia(tmp_path):
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", columnas_extra=("Observaciones",)))
    assert rep.es_valido
    assert "columna_nueva" in codigos(rep, "advertencia")


def test_formato_sin_clasificacion(tmp_path):
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", clasificacion=False))
    assert rep.es_valido and not rep.trae_clasificacion


def test_falta_columna_obligatoria(tmp_path):
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", quitar_columnas=("Límite de asignados",)))
    assert "faltan_columnas" in codigos(rep, "error")


def test_falta_una_hoja(tmp_path):
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", hojas=("Clase",)))
    assert rep.es_valido
    assert "falta_hoja" in codigos(rep, "advertencia")


def test_sin_hojas_reconocidas(tmp_path):
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", hojas=("Hoja1",)))
    assert "sin_hojas" in codigos(rep, "error")


def test_nombre_de_hoja_con_otra_capitalizacion(tmp_path):
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", hojas=("CLASE ", "laboratorio")))
    assert rep.es_valido and set(rep.hojas_leidas) == {"Clase", "Laboratorio"}


def test_totales_del_pie_no_cuadran(tmp_path):
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", total_pie={"total": 12345}))
    assert "control_total" in codigos(rep, "error")


def test_regulares_del_pie_no_cuadran(tmp_path):
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", total_pie={"reg": 1}))
    assert "control_regulares" in codigos(rep, "error")


def test_sin_pie_de_control(tmp_path):
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", sin_pie=True))
    assert rep.es_valido
    assert "sin_pie_control" in codigos(rep, "advertencia")


@pytest.mark.parametrize("total,codigo", [("muchos", "valor_no_numerico"), (-5, "valor_negativo")])
def test_cantidades_invalidas(tmp_path, total, codigo):
    secs = [dict(SECCIONES_CLASE[0], total=total)]
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", clase=secs, laboratorio=[], sin_pie=True))
    assert codigo in codigos(rep, "error")


def test_total_distinto_de_suma_de_procedencia(tmp_path):
    secs = [dict(SECCIONES_CLASE[0], total=500)]
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", clase=secs, laboratorio=[]))
    assert "suma_procedencia" in codigos(rep, "advertencia")


def test_hora_invalida(tmp_path):
    secs = [dict(SECCIONES_CLASE[0], hi="25:99", hf="07:00")]
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", clase=secs, laboratorio=[]))
    assert rep.es_valido and "hora_invalida" in codigos(rep, "advertencia")


def test_seccion_repetida(tmp_path):
    secs = [SECCIONES_CLASE[0], SECCIONES_CLASE[0]]
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", clase=secs, laboratorio=[]))
    assert "seccion_duplicada" in codigos(rep, "advertencia")


def test_horas_como_objetos_time_y_codigo_numerico(tmp_path):
    ruta = construir_reporte(tmp_path / "a.xlsx", laboratorio=[])
    wb = load_workbook(ruta)
    ws = wb["Clase"]
    ws["G8"], ws["B8"] = time(7, 10), 770
    wb.save(ruta)
    rep = lector.leer_reporte(ruta)
    s = rep.secciones[0]
    assert rep.es_valido and s["hora_inicio"] == "07:10" and s["codigo"] == "0770"


# ═══════════════════════════ membrete ═══════════════════════════
def test_sin_membrete(tmp_path):
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", membrete="REPORTE CUALQUIERA"))
    assert "sin_membrete" in codigos(rep, "error")


def test_membrete_ilegible(tmp_path):
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", membrete="TOTAL DE ASIGNADOS ALGO"))
    assert "membrete_ilegible" in codigos(rep, "error")


def test_periodo_fuera_de_catalogo(tmp_path):
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", periodo="CURSO ESPECIAL"))
    assert rep.es_valido and rep.periodo is None
    assert "periodo_desconocido" in codigos(rep, "advertencia")


def test_periodo_con_tildes(tmp_path):
    rep = lector.leer_reporte(construir_reporte(
        tmp_path / "a.xlsx", periodo="EXAMEN UBICACIÓN IDIOMA TÉCNICO - SEGUNDO SEMESTRE"))
    assert rep.es_valido and rep.periodo.codigo == "21"


def test_corte_en_el_futuro(tmp_path):
    rep = lector.leer_reporte(construir_reporte(tmp_path / "a.xlsx", corte=datetime(2099, 1, 1)))
    assert rep.es_valido and "corte_futuro" in codigos(rep, "advertencia")


def test_hojas_de_reportes_distintos(tmp_path):
    ruta = construir_reporte(tmp_path / "a.xlsx")
    wb = load_workbook(ruta)
    wb["Laboratorio"]["A4"] = "TOTAL DE ASIGNADOS SEGUNDO SEMESTRE - 2026, FECHA/HORA: 01/10/2026  08:00"
    wb.save(ruta)
    assert "membretes_distintos" in codigos(lector.leer_reporte(ruta), "error")


# ═══════════════════════════ archivo (bytes) ═══════════════════════════
def test_archivo_vacio(tmp_path):
    (tmp_path / "a.xlsx").write_bytes(b"")
    assert "archivo_vacio_bytes" in codigos(lector.leer_reporte(tmp_path / "a.xlsx"))


def test_xls_antiguo_o_con_contrasena(tmp_path):
    (tmp_path / "a.xls").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 2048)
    assert "formato_xls_o_protegido" in codigos(lector.leer_reporte(tmp_path / "a.xls"))


@pytest.mark.parametrize("contenido", [b"Codigo;Nombre;Total\n0770;IPC;10\n", b"<html><table></table></html>",
                                       b"\xef\xbb\xbf<?xml version='1.0'?><Workbook/>"])
def test_texto_renombrado_como_xlsx(tmp_path, contenido):
    (tmp_path / "a.xlsx").write_bytes(contenido)
    assert "no_es_xlsx" in codigos(lector.leer_reporte(tmp_path / "a.xlsx"))


def test_zip_que_no_es_excel(tmp_path):
    with zipfile.ZipFile(tmp_path / "a.xlsx", "w") as z:
        z.writestr("hola.txt", "hola")
    assert "zip_no_xlsx" in codigos(lector.leer_reporte(tmp_path / "a.xlsx"))


def test_descarga_interrumpida(tmp_path):
    datos = REAL_2026.read_bytes()
    (tmp_path / "a.xlsx").write_bytes(datos[: len(datos) // 2])
    assert codigos(lector.leer_reporte(tmp_path / "a.xlsx"), "error") & {"zip_danado", "libro_ilegible"}


def test_zip_bomba(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "MAX_DESCOMPRIMIDO_BYTES", 1024 * 1024)
    ruta = construir_reporte(tmp_path / "a.xlsx")
    with zipfile.ZipFile(ruta, "a", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/media/relleno.bin", b"\0" * (5 * 1024 * 1024))
    assert "zip_bomba" in codigos(lector.leer_reporte(ruta), "error")


def test_archivo_demasiado_grande(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "MAX_BYTES", 1000)
    assert "archivo_muy_grande" in codigos(lector.leer_reporte(REAL_2026), "error")


def test_libro_con_macros(tmp_path):
    ruta = construir_reporte(tmp_path / "a.xlsm")
    with zipfile.ZipFile(ruta, "a") as z:
        z.writestr("xl/vbaProject.bin", b"macro")
    rep = lector.leer_reporte(ruta)
    assert rep.es_valido and "con_macros" in codigos(rep, "advertencia")


def test_extension_equivocada_pero_contenido_valido(tmp_path):
    ruta = construir_reporte(tmp_path / "a.xlsx").rename(tmp_path / "reporte.xls")
    rep = lector.leer_reporte(ruta)
    assert rep.es_valido and "extension_inesperada" in codigos(rep, "advertencia")


def test_xml_malicioso_no_se_expande(tmp_path):
    """Billion laughs dentro de una hoja: debe rechazarse sin colgar el servidor."""
    ruta = construir_reporte(tmp_path / "a.xlsx")
    bomba = ('<?xml version="1.0"?><!DOCTYPE l [<!ENTITY a "aaaaaaaaaa">'
             + "".join(f'<!ENTITY {chr(98 + i)} "{("&" + chr(97 + i) + ";") * 10}">' for i in range(8))
             + ']><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
             '<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&i;</t></is></c></row></sheetData></worksheet>')
    salida = io.BytesIO()
    with zipfile.ZipFile(ruta) as zin, zipfile.ZipFile(salida, "w") as zout:
        for item in zin.infolist():
            datos = bomba.encode() if item.filename == "xl/worksheets/sheet1.xml" else zin.read(item)
            zout.writestr(item, datos)
    ruta.write_bytes(salida.getvalue())
    rep = lector.leer_reporte(ruta)
    assert not rep.es_valido

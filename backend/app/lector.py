"""
Lectura y validación del reporte «TOTAL DE ASIGNADOS» (.xlsx de Centro de Cálculo).

Este módulo no toca la base de datos: recibe una ruta y devuelve un `Reporte` con
las secciones normalizadas y la lista de mensajes de validación. La decisión de
cargar o rechazar la toma `ingesta.py`.

Variaciones reales del formato que este lector soporta (todas vistas en archivos
reales, ver backend/tests/fixtures/):

* 30, 32 o 33 columnas: algunas emisiones agregan Escuela / Area / Subarea y una
  o dos parejas «Nombre miembro / Tipo miembro». Por eso las columnas se
  resuelven **por nombre**, no por posición.
* «Nombre» aparece dos veces: la primera es el curso y la segunda el catedrático.
* Una sección puede ocupar varias filas (un bloque de horario por fila); las de
  continuación dejan vacía la columna Total.
* Secciones cerradas: edificio ET, salón CERRADO, catedrático STAFF.
* Reportes vacíos: sin fila de encabezado, solo el pie con «TOTAL: 0».
* Pie con totales de control («TOTAL REG:» y «TOTAL:»), que se usan como suma de
  verificación de toda la lectura.
* La Ñ llega corrupta como «\\» en algunos nombres (DISE\\O).

Niveles de mensaje:
  error        -> el archivo NO se carga.
  advertencia  -> se carga, pero el usuario debe saberlo.
  info         -> contexto (p. ej. una hoja vacía en un reporte que sí trae datos).
"""

from __future__ import annotations

import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from . import catalogos, config
from .catalogos import normalizar

HOJAS_ESPERADAS = {"CLASE": "Clase", "LABORATORIO": "Laboratorio"}
EXTENSIONES = {".xlsx", ".xlsm"}

# Columnas por nombre normalizado -> clave interna. "NOMBRE" se resuelve aparte.
COLUMNAS = {
    "NO": "no",
    "CODIGO": "codigo",
    "SECCION": "seccion",
    "EDIFICIO": "edificio",
    "SALON": "salon",
    "HORA INICIO": "hora_inicio",
    "HORA FINAL": "hora_fin",
    "L": "L", "MA": "Ma", "MI": "Mi", "J": "J", "V": "V", "S": "S", "D": "D",
    "HORAS DIA": "horas_dia",
    "HORAS SEMANA": "horas_semana",
    "REGULARES": "regulares",
    "OTRA FAC.": "otra_fac",
    "CENTROS REG.": "centros_reg",
    "OTRA UNIV.": "otra_univ",
    "NO INSC.": "no_insc",
    "EST. CIERRE": "est_cierre",
    "OTRO CASO": "otro_caso",
    "TOTAL": "total",
    "LIMITE DE ASIGNADOS": "limite",
    "LIMITE MINIMO DE ASIGNADOS": "limite_min",
    "ESCUELA": "escuela",
    "AREA": "area",
    "SUBAREA": "subarea",
}
OBLIGATORIAS = [
    "codigo", "curso", "seccion", "edificio", "salon", "hora_inicio", "hora_fin",
    "L", "Ma", "Mi", "J", "V", "S", "D", "docente", "total", "limite",
]
PROCEDENCIA = [c for c, _ in config.COLUMNAS_PROCEDENCIA]

# Marca que dejan algunas herramientas de exportación cuando cortan una tabla larga
# («... (truncated at 1000 rows)»). Un archivo así está incompleto.
RE_RECORTE = re.compile(r"TRUNCAT|RECORTAD|FILAS OMITIDAS|ROWS OMITTED|MORE ROWS")
TOLERANCIA_MEMBRETE_MIN = 60

RE_MEMBRETE = re.compile(
    r"TOTAL DE ASIGNADOS\s+(?P<periodo>.+?)\s*-\s*(?P<anio>\d{4})\s*,?\s*"
    r"FECHA\s*/\s*HORA\s*:\s*(?P<d>\d{1,2})/(?P<m>\d{1,2})/(?P<y>\d{4})"
    r"(?:\s+(?P<h>\d{1,2}):(?P<mi>\d{2}))?",
    re.I,
)


# ═══════════════════════════════════════════════════════════════════════════
# Resultado
# ═══════════════════════════════════════════════════════════════════════════
@dataclass
class Mensaje:
    nivel: str          # error | advertencia | info
    codigo: str         # identificador estable, útil para pruebas y para la UI
    texto: str
    hoja: str | None = None
    fila: int | None = None

    def a_dict(self) -> dict:
        return {k: v for k, v in self.__dict__.items() if v is not None}


@dataclass
class Reporte:
    ruta: Path
    tipo_documento: str = "ASIGNADOS_CONDENSADO"
    encabezado: str | None = None
    periodo_nombre: str | None = None
    periodo: catalogos.Periodo | None = None
    anio: int | None = None
    fecha_corte: datetime | None = None
    secciones: list[dict] = field(default_factory=list)
    mensajes: list[Mensaje] = field(default_factory=list)
    hojas_leidas: list[str] = field(default_factory=list)
    hojas_vacias: list[str] = field(default_factory=list)
    columnas_extra: list[str] = field(default_factory=list)
    trae_clasificacion: bool = False
    control: dict = field(default_factory=dict)   # hoja -> {"total_reg", "total"}

    # -- atajos ---------------------------------------------------------------
    def error(self, codigo, texto, hoja=None, fila=None):
        self.mensajes.append(Mensaje("error", codigo, texto, hoja, fila))

    def advertencia(self, codigo, texto, hoja=None, fila=None):
        self.mensajes.append(Mensaje("advertencia", codigo, texto, hoja, fila))

    def info(self, codigo, texto, hoja=None, fila=None):
        self.mensajes.append(Mensaje("info", codigo, texto, hoja, fila))

    @property
    def errores(self) -> list[Mensaje]:
        return [m for m in self.mensajes if m.nivel == "error"]

    @property
    def advertencias(self) -> list[Mensaje]:
        return [m for m in self.mensajes if m.nivel == "advertencia"]

    @property
    def es_valido(self) -> bool:
        return not self.errores

    @property
    def vacio(self) -> bool:
        """Reporte bien formado pero sin ninguna sección (p. ej. finales 2023)."""
        return self.es_valido and not self.secciones and bool(self.hojas_vacias)

    @property
    def total_asignados(self) -> int:
        return sum(s["total"] for s in self.secciones)

    @property
    def descripcion(self) -> str:
        p = self.periodo.nombre if self.periodo else (self.periodo_nombre or "período desconocido")
        corte = self.fecha_corte.strftime("%d/%m/%Y %H:%M") if self.fecha_corte else "s/f"
        return f"{p} {self.anio or ''} · corte {corte}".strip()


# ═══════════════════════════════════════════════════════════════════════════
# Normalización de celdas
# ═══════════════════════════════════════════════════════════════════════════
def limpiar_texto(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    s = s.replace("\\", "Ñ")             # DISE\O -> DISEÑO (Ñ corrupta del origen)
    s = re.sub(r"\s+", " ", s)
    s = unicodedata.normalize("NFC", s)
    return s or None


def limpiar_codigo(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    if not s:
        return None
    return s.zfill(4) if s.isdigit() else s.upper()


def limpiar_salon(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s.upper() if s else None


def limpiar_hora(v) -> str | None:
    if v is None:
        return None
    if hasattr(v, "strftime"):
        return v.strftime("%H:%M")
    s = str(v).strip()
    m = re.match(r"^(\d{1,2})[:.](\d{2})", s)
    if m and int(m.group(1)) < 24 and int(m.group(2)) < 60:
        return f"{int(m.group(1)):02d}:{m.group(2)}"
    return None


class ValorInvalido(ValueError):
    pass


def numero(v, entero: bool = True):
    """Convierte una celda numérica. Vacío -> 0. Texto no numérico -> ValorInvalido."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return 0
    if isinstance(v, bool):
        raise ValorInvalido(v)
    if isinstance(v, (int, float)):
        return int(v) if entero else float(v)
    s = str(v).strip().replace(",", "")
    try:
        f = float(s)
    except ValueError as exc:
        raise ValorInvalido(v) from exc
    return int(f) if entero else f


# ═══════════════════════════════════════════════════════════════════════════
# 1. Inspección del archivo (antes de abrirlo con openpyxl)
# ═══════════════════════════════════════════════════════════════════════════
def inspeccionar_archivo(ruta: Path, nombre_original: str | None = None) -> list[Mensaje]:
    """Validaciones de bajo nivel: tamaño, tipo real del contenido, zip bomb."""
    msgs: list[Mensaje] = []
    nombre = nombre_original or ruta.name
    ext = Path(nombre).suffix.lower()

    try:
        tam = ruta.stat().st_size
    except OSError as exc:
        return [Mensaje("error", "archivo_ilegible", f"No se pudo leer el archivo: {exc}")]
    if tam == 0:
        return [Mensaje("error", "archivo_vacio_bytes", "El archivo está vacío (0 bytes).")]
    if tam > config.MAX_BYTES:
        return [Mensaje(
            "error", "archivo_muy_grande",
            f"El archivo pesa {tam / 1048576:.1f} MB y el máximo permitido es "
            f"{config.MAX_MB:g} MB. Un reporte de asignados real pesa menos de 1 MB.",
        )]

    try:
        with ruta.open("rb") as f:
            cabeza = f.read(8)
    except PermissionError:
        return [Mensaje(
            "error", "archivo_ilegible",
            "El sistema no tiene permiso para leer el archivo. En el equipo servidor, dé permiso de "
            "lectura al archivo (en Linux: chmod a+r) y vuelva a copiarlo a la carpeta de entrada.")]
    except OSError as exc:
        return [Mensaje("error", "archivo_ilegible", f"No se pudo leer el archivo: {exc}")]

    if cabeza.startswith(b"\xd0\xcf\x11\xe0"):
        return [Mensaje(
            "error", "formato_xls_o_protegido",
            "El archivo es un Excel antiguo (.xls) o está protegido con contraseña. "
            "Ábralo en Excel y use «Guardar como» → «Libro de Excel (.xlsx)» sin contraseña.",
        )]
    if not cabeza.startswith(b"PK"):
        muestra = cabeza.lstrip(b"\xef\xbb\xbf").lstrip()[:1]
        tipo = "HTML o XML" if muestra == b"<" else "texto (CSV u otro)"
        return [Mensaje(
            "error", "no_es_xlsx",
            f"El contenido no es un libro de Excel (.xlsx); parece {tipo}. "
            "Descargue de nuevo el reporte «Contadores asignados (Condensado)» en formato Excel.",
        )]

    try:
        with zipfile.ZipFile(ruta) as z:
            nombres = set(z.namelist())
            descomprimido = sum(i.file_size for i in z.infolist())
    except zipfile.BadZipFile:
        return [Mensaje("error", "zip_danado",
                        "El archivo está dañado o incompleto (¿se interrumpió la descarga?).")]

    if "xl/workbook.xml" not in nombres:
        return [Mensaje("error", "zip_no_xlsx",
                        "Es un archivo comprimido (.zip), no un libro de Excel. Si trae reportes adentro, "
                        "súbalo en «Cargar reportes» o cópielo a la carpeta de entrada: cada .xlsx "
                        "de adentro se carga por separado.")]
    if descomprimido > config.MAX_DESCOMPRIMIDO_BYTES:
        return [Mensaje(
            "error", "zip_bomba",
            f"El contenido descomprimido ({descomprimido / 1048576:.0f} MB) excede el límite "
            "de seguridad. Un reporte de asignados no es tan grande.",
        )]
    if "xl/vbaProject.bin" in nombres:
        msgs.append(Mensaje("advertencia", "con_macros",
                            "El libro contiene macros. Se ignoran: solo se leen los datos."))
    if ext not in EXTENSIONES:
        msgs.append(Mensaje(
            "advertencia", "extension_inesperada",
            f"La extensión «{ext or '(ninguna)'}» no es .xlsx, pero el contenido sí es un "
            "libro de Excel válido; se procesa por su contenido.",
        ))
    return msgs


# ═══════════════════════════════════════════════════════════════════════════
# 2. Lectura del libro
# ═══════════════════════════════════════════════════════════════════════════
def _leer_membrete(filas: list[tuple]) -> tuple[str | None, re.Match | None]:
    for fila in filas[:12]:
        for celda in fila[:3]:
            txt = limpiar_texto(celda)
            if txt and "TOTAL DE ASIGNADOS" in normalizar(txt):
                return txt, RE_MEMBRETE.search(normalizar(txt))
    return None, None


def _localizar_encabezado(filas: list[tuple]) -> int | None:
    for i, fila in enumerate(filas[:30]):
        if fila and normalizar(fila[0]) == "NO" and any(normalizar(c) == "CODIGO" for c in fila[:5]):
            return i
    return None


def _mapear_columnas(rep: Reporte, hoja: str, encabezado: tuple) -> dict | None:
    mapa: dict[str, int] = {}
    miembros: list[tuple[int, int | None]] = []
    nombres_vistos = 0
    for i, celda in enumerate(encabezado):
        n = normalizar(celda)
        if not n:
            continue
        if n == "NOMBRE":
            mapa["curso" if nombres_vistos == 0 else "docente"] = i
            nombres_vistos += 1
        elif n == "NOMBRE MIEMBRO":
            siguiente = normalizar(encabezado[i + 1]) if i + 1 < len(encabezado) else ""
            miembros.append((i, i + 1 if siguiente == "TIPO MIEMBRO" else None))
        elif n == "TIPO MIEMBRO":
            continue
        elif n in COLUMNAS:
            mapa.setdefault(COLUMNAS[n], i)
        else:
            original = limpiar_texto(celda)
            if original not in rep.columnas_extra:
                rep.columnas_extra.append(original)
                rep.advertencia(
                    "columna_nueva",
                    f"Columna no reconocida «{original}»: se ignora. Si el reporte cambió de "
                    "formato, revise si esta columna debería incorporarse.", hoja)

    faltan = [c for c in OBLIGATORIAS if c not in mapa]
    if faltan:
        legibles = {"curso": "Nombre (curso)", "docente": "Nombre (catedrático)"}
        rep.error(
            "faltan_columnas",
            "Faltan columnas obligatorias: " + ", ".join(legibles.get(c, c) for c in faltan) +
            ". El archivo no tiene el formato del reporte de asignados.", hoja)
        return None
    mapa["_miembros"] = miembros  # type: ignore[assignment]
    if {"escuela", "area"} <= mapa.keys():
        rep.trae_clasificacion = True
    return mapa


def _leer_pie(fila: tuple) -> dict | None:
    """Detecta la fila «TOTAL REG: n … TOTAL: n» y devuelve sus valores."""
    vals = {}
    for i, celda in enumerate(fila):
        n = normalizar(celda).rstrip(":").strip()
        if n in ("TOTAL REG", "TOTAL") and i + 1 < len(fila):
            try:
                vals["total_reg" if n == "TOTAL REG" else "total"] = numero(fila[i + 1])
            except ValorInvalido:
                vals["total_reg" if n == "TOTAL REG" else "total"] = None
    return vals or None


def _parsear_hoja(rep: Reporte, ws, tipo: str) -> None:
    filas = list(ws.iter_rows(values_only=True))
    hoja = ws.title

    for n_fila, fila in enumerate(filas, start=1):
        for celda_txt in (fila or ())[:3]:
            if isinstance(celda_txt, str) and RE_RECORTE.search(normalizar(celda_txt)):
                rep.error("archivo_recortado",
                          f"La hoja «{hoja}» está recortada: en la fila {n_fila} dice "
                          f"«{limpiar_texto(celda_txt)}». Faltan secciones, así que el reporte "
                          "está incompleto. Descárguelo de nuevo desde Centro de Cálculo, sin "
                          "pasarlo por otra herramienta.", hoja, n_fila)
                return

    i_enc = _localizar_encabezado(filas)
    pie = None
    for fila in reversed(filas):
        if fila and any(v is not None for v in fila):
            pie = _leer_pie(fila)
            break

    if i_enc is None:
        if pie is not None and (pie.get("total") or 0) == 0:
            rep.hojas_vacias.append(tipo)
            rep.control[tipo] = pie
            rep.info("hoja_vacia", f"La hoja «{hoja}» no tiene secciones (TOTAL: 0).", hoja)
        else:
            rep.error("sin_encabezado",
                      f"No se encontró la fila de encabezado en la hoja «{hoja}» "
                      "(se busca una fila que empiece con «No» y «Codigo»).", hoja)
        return

    mapa = _mapear_columnas(rep, hoja, filas[i_enc])
    if mapa is None:
        return
    rep.hojas_leidas.append(tipo)
    col = mapa
    miembros_cols = mapa.pop("_miembros")  # type: ignore[arg-type]

    def celda(fila, clave):
        i = col.get(clave)
        return fila[i] if i is not None and i < len(fila) else None

    actual: dict | None = None
    claves_vistas: dict[tuple, int] = {}
    avisos = {"suma_procedencia": [], "hora": [], "codigo": [], "huerfana": [], "duplicada": []}
    pie_hoja = None

    for n_fila, fila in enumerate(filas[i_enc + 1:], start=i_enc + 2):
        if not fila or all(v is None for v in fila):
            continue
        codigo = limpiar_codigo(celda(fila, "codigo"))
        if not codigo:
            p = _leer_pie(fila)
            if p:
                pie_hoja = p
            else:
                avisos["huerfana"].append(n_fila)
            continue

        seccion = limpiar_texto(celda(fila, "seccion")) or ""
        h_ini = limpiar_hora(celda(fila, "hora_inicio"))
        h_fin = limpiar_hora(celda(fila, "hora_fin"))
        dias = [d for d in config.DIAS if limpiar_texto(celda(fila, d))]
        trae_total = celda(fila, "total") is not None

        if (celda(fila, "hora_inicio") is not None and not h_ini) or \
           (celda(fila, "hora_fin") is not None and not h_fin) or \
           (h_ini and h_fin and h_fin <= h_ini):
            avisos["hora"].append(n_fila)

        if not trae_total and actual is not None and \
                actual["codigo"] == codigo and actual["seccion"] == seccion:
            for d in dias:
                actual["bloques"].append({"dia": d, "hora_inicio": h_ini, "hora_fin": h_fin})
            try:
                actual["horas_semana"] += numero(celda(fila, "horas_semana"), entero=False)
            except ValorInvalido:
                pass
            continue

        if not re.fullmatch(r"\d{4}", codigo):
            avisos["codigo"].append(f"{codigo} (fila {n_fila})")

        try:
            valores = {
                k: numero(celda(fila, k)) for k in
                ["total", "limite", "limite_min", *PROCEDENCIA] if k in col
            }
            horas_dia = numero(celda(fila, "horas_dia"), entero=False)
            horas_semana = numero(celda(fila, "horas_semana"), entero=False)
        except ValorInvalido as exc:
            rep.error("valor_no_numerico",
                      f"Valor no numérico «{exc.args[0]}» en una columna de cantidades.",
                      hoja, n_fila)
            continue
        if any(v < 0 for v in valores.values()):
            rep.error("valor_negativo", "Hay cantidades negativas.", hoja, n_fila)
            continue
        if not trae_total:
            rep.advertencia("fila_sin_total",
                            f"La sección {codigo}-{seccion} no trae Total; se carga con 0.",
                            hoja, n_fila)

        total = valores.get("total", 0)
        proc_presentes = [k for k in PROCEDENCIA if k in col]
        if proc_presentes and trae_total:
            suma = sum(valores.get(k, 0) for k in proc_presentes)
            if suma != total:
                avisos["suma_procedencia"].append(f"{codigo}-{seccion} (fila {n_fila})")

        clave = (codigo, seccion)
        if clave in claves_vistas:
            avisos["duplicada"].append(f"{codigo}-{seccion} (filas {claves_vistas[clave]} y {n_fila})")
        claves_vistas[clave] = n_fila

        edificio = limpiar_texto(celda(fila, "edificio"))
        salon = limpiar_salon(celda(fila, "salon"))
        docente = limpiar_texto(celda(fila, "docente"))
        cerrada = (salon or "") == config.SALON_CERRADO
        sin_docente = normalizar(docente) in config.DOCENTES_SIN_ASIGNAR

        def clasif(k):
            v = limpiar_texto(celda(fila, k))
            return None if normalizar(v) in ("", "SIN SUBAREA", "SIN AREA") else normalizar(v)

        miembros, vistos = [], set()
        for i_nom, i_tipo in miembros_cols:
            nom = limpiar_texto(fila[i_nom]) if i_nom < len(fila) else None
            tip = limpiar_texto(fila[i_tipo]) if i_tipo is not None and i_tipo < len(fila) else None
            if nom and normalizar(nom) not in vistos:
                vistos.add(normalizar(nom))
                miembros.append({"nombre": nom, "tipo": (tip or "SIN TIPO").upper()})

        actual = {
            "tipo": tipo,
            "codigo": codigo,
            "curso": limpiar_texto(celda(fila, "curso")) or "(sin nombre)",
            "seccion": seccion,
            "edificio": edificio,
            "salon": salon,
            "modalidad": config.modalidad_de(edificio, salon),
            "jornada": config.jornada_de(h_ini),
            "hora_inicio": h_ini,
            "hora_fin": h_fin,
            "horas_dia": horas_dia,
            "horas_semana": horas_semana,
            "docente": None if sin_docente else docente,
            "sin_docente": sin_docente and not cerrada,
            "cerrada": cerrada,
            **{k: valores.get(k, 0) for k in PROCEDENCIA},
            "total": total,
            "limite": valores.get("limite", 0),
            "limite_min": valores.get("limite_min", 0),
            "escuela": clasif("escuela"),
            "area": clasif("area"),
            "subarea": clasif("subarea"),
            "fila_origen": n_fila,
            "bloques": [{"dia": d, "hora_inicio": h_ini, "hora_fin": h_fin} for d in dias],
            "miembros": miembros,
        }
        actual["estado"] = config.estado_de(total, actual["limite"], actual["limite_min"], cerrada)
        rep.secciones.append(actual)

    # --- avisos agrupados (uno por tipo, con ejemplos) -------------------------
    def resumen(lista, n=4):
        return ", ".join(map(str, lista[:n])) + (f" y {len(lista) - n} más" if len(lista) > n else "")

    if avisos["suma_procedencia"]:
        rep.advertencia("suma_procedencia",
                        f"{len(avisos['suma_procedencia'])} secciones donde Total no es la suma de "
                        f"las columnas de procedencia: {resumen(avisos['suma_procedencia'])}.", hoja)
    if avisos["hora"]:
        rep.advertencia("hora_invalida",
                        f"{len(avisos['hora'])} filas con hora inválida o final antes del inicio "
                        f"(filas {resumen(avisos['hora'])}).", hoja)
    if avisos["codigo"]:
        rep.advertencia("codigo_inusual",
                        f"Códigos de curso con formato inusual: {resumen(avisos['codigo'])}.", hoja)
    if avisos["huerfana"]:
        rep.advertencia("filas_sin_codigo",
                        f"{len(avisos['huerfana'])} filas sin código de curso se ignoraron "
                        f"(filas {resumen(avisos['huerfana'])}).", hoja)
    if avisos["duplicada"]:
        rep.advertencia("seccion_duplicada",
                        f"Secciones repetidas con su propio total: {resumen(avisos['duplicada'])}.",
                        hoja)

    # --- suma de verificación contra el pie del reporte ------------------------
    propias = [s for s in rep.secciones if s["tipo"] == tipo]
    suma_total = sum(s["total"] for s in propias)
    suma_reg = sum(s["regulares"] for s in propias)
    if pie_hoja is None:
        rep.advertencia("sin_pie_control",
                        f"La hoja «{hoja}» no trae la fila de totales de control (TOTAL REG / "
                        "TOTAL); no se pudo verificar la suma.", hoja)
        return
    rep.control[tipo] = pie_hoja
    if pie_hoja.get("total") is not None and pie_hoja["total"] != suma_total:
        rep.error("control_total",
                  f"La suma de asignados leída ({suma_total:,}) no coincide con el TOTAL del pie "
                  f"del reporte ({pie_hoja['total']:,}). El archivo pudo haber sido editado o "
                  "estar incompleto.", hoja)
    if "regulares" in col and pie_hoja.get("total_reg") is not None \
            and pie_hoja["total_reg"] != suma_reg:
        rep.error("control_regulares",
                  f"La suma de regulares leída ({suma_reg:,}) no coincide con TOTAL REG del pie "
                  f"({pie_hoja['total_reg']:,}).", hoja)


def leer_reporte(ruta: Path, nombre_original: str | None = None) -> Reporte:
    ruta = Path(ruta)
    rep = Reporte(ruta=ruta)
    rep.mensajes.extend(inspeccionar_archivo(ruta, nombre_original))
    if rep.errores:
        return rep

    # Se abre como flujo binario: openpyxl rechaza por EXTENSIÓN las rutas que no
    # terminan en .xlsx (las subidas llegan como temporales), y el tipo real ya se
    # verificó por contenido en inspeccionar_archivo.
    flujo = ruta.open("rb")
    try:
        wb = load_workbook(flujo, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl lanza varios tipos
        flujo.close()
        rep.error("libro_ilegible",
                  f"El libro de Excel está dañado o no se puede leer ({type(exc).__name__}). "
                  "Descargue el reporte de nuevo.")
        return rep

    try:
        hojas = {normalizar(n): n for n in wb.sheetnames}
        encontradas = {k: hojas[k] for k in HOJAS_ESPERADAS if k in hojas}
        if not encontradas:
            rep.error("sin_hojas",
                      "No se encontraron las hojas «Clase» ni «Laboratorio». Hojas del archivo: "
                      + ", ".join(wb.sheetnames) + ".")
            return rep
        for k, legible in HOJAS_ESPERADAS.items():
            if k not in encontradas:
                rep.advertencia("falta_hoja",
                                f"El archivo no trae la hoja «{legible}»; se carga solo lo demás.")

        membretes = {}
        for clave, nombre_hoja in encontradas.items():
            ws = wb[nombre_hoja]
            filas_cabeza = list(ws.iter_rows(min_row=1, max_row=12, values_only=True))
            membretes[clave] = _leer_membrete(filas_cabeza)
            _parsear_hoja(rep, ws, HOJAS_ESPERADAS[clave])
    except Exception as exc:  # noqa: BLE001 — libro con XML inválido a mitad de lectura
        rep.error("libro_ilegible",
                  f"El libro de Excel tiene contenido dañado ({type(exc).__name__}). "
                  "Descargue el reporte de nuevo.")
        return rep
    finally:
        wb.close()
        flujo.close()

    _resolver_membrete(rep, membretes)
    if rep.es_valido and not rep.secciones and not rep.hojas_vacias:
        rep.error("sin_datos", "El archivo no contiene secciones legibles.")
    return rep


def _momento(m: re.Match) -> datetime | None:
    try:
        return datetime(int(m.group("y")), int(m.group("m")), int(m.group("d")),
                        int(m.group("h") or 0), int(m.group("mi") or 0))
    except ValueError:
        return None


def _unificar_membretes(rep: Reporte, membretes: dict, textos: set) -> set:
    """Centro de Cálculo genera cada hoja por separado: si la descarga cruza un cambio de
    minuto, «Clase» dice 11:38 y «Laboratorio» 11:39. Es el mismo reporte cuando período y
    año coinciden y las horas difieren poco; se toma la hora más reciente como corte."""
    leidos = [(t, m) for t, m in membretes.values() if t and m]
    if len(leidos) != len([t for t, _ in membretes.values() if t]):
        return textos
    claves = {(normalizar(m.group("periodo")).strip(), m.group("anio")) for _, m in leidos}
    momentos = [_momento(m) for _, m in leidos]
    if len(claves) != 1 or any(x is None for x in momentos):
        return textos
    diferencia = (max(momentos) - min(momentos)).total_seconds() / 60
    if diferencia > TOLERANCIA_MEMBRETE_MIN:
        return textos
    texto = max(leidos, key=lambda tm: _momento(tm[1]))[0]
    rep.info("membretes_minutos",
             f"Las hojas se generaron con {int(diferencia)} min de diferencia "
             f"({min(momentos):%H:%M} y {max(momentos):%H:%M}); se toma "
             f"{max(momentos):%d/%m/%Y %H:%M} como hora del corte.")
    return {normalizar(texto)}


def _resolver_membrete(rep: Reporte, membretes: dict) -> None:
    textos = {normalizar(t) for t, _ in membretes.values() if t}
    if len(textos) > 1:
        textos = _unificar_membretes(rep, membretes, textos)
    if not textos:
        rep.error("sin_membrete",
                  "No se encontró la línea «TOTAL DE ASIGNADOS <PERÍODO> - <AÑO>, FECHA/HORA: …». "
                  "Sin ella no se puede saber a qué período ni a qué corte pertenece el archivo.")
        return
    if len(textos) > 1:
        rep.error("membretes_distintos",
                  "Las hojas del archivo dicen pertenecer a reportes distintos: "
                  + " | ".join(sorted(textos)))
        return

    texto, m = next((t, m) for t, m in membretes.values() if t and normalizar(t) in textos)
    rep.encabezado = texto
    if not m:
        rep.error("membrete_ilegible",
                  f"No se pudo interpretar la línea del reporte: «{texto}». Se espera "
                  "«TOTAL DE ASIGNADOS <PERÍODO> - <AÑO>, FECHA/HORA: dd/mm/aaaa hh:mm».")
        return

    rep.periodo_nombre = m.group("periodo").strip()
    rep.anio = int(m.group("anio"))
    rep.periodo = catalogos.periodo_por_nombre(rep.periodo_nombre)
    if rep.periodo is None:
        rep.advertencia("periodo_desconocido",
                        f"El período «{rep.periodo_nombre}» no está en el catálogo de 21 períodos "
                        "de Centro de Cálculo. Se carga con ese nombre, al final del año.")
    try:
        rep.fecha_corte = datetime(
            int(m.group("y")), int(m.group("m")), int(m.group("d")),
            int(m.group("h") or 0), int(m.group("mi") or 0),
        )
    except ValueError:
        rep.error("fecha_corte_invalida", f"La fecha/hora del reporte no es válida: «{texto}».")
        return

    ahora = datetime.now(ZoneInfo(config.ZONA_HORARIA)).replace(tzinfo=None)
    if rep.fecha_corte > ahora + timedelta(days=1):
        rep.advertencia("corte_futuro",
                        f"La fecha de corte ({rep.fecha_corte:%d/%m/%Y %H:%M}) está en el futuro. "
                        "Revise el reloj del equipo donde se descargó el reporte.")
    if not (2000 <= rep.anio <= rep.fecha_corte.year + 1):
        rep.advertencia("anio_inusual",
                        f"El año académico {rep.anio} no es coherente con la fecha de corte "
                        f"{rep.fecha_corte:%d/%m/%Y}.")

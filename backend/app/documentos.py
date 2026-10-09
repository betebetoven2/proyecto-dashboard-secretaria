"""
Registro de tipos de documento.

El sistema recibe archivos de Centro de Cálculo que pueden ser de distintos
TIPOS (hoy solo «Contadores asignados (Condensado)») y, dentro de un tipo, de
distintos FORMATOS (30/32/33 columnas) y de distintos PERÍODOS (21 tipos de
período: semestre, vacaciones, finales, retrasadas…).

Esas tres dimensiones son independientes y se manejan en capas distintas:

    tipo de documento  -> este módulo: decide qué lector usar, o rechaza
    formato            -> el lector del tipo (columnas por nombre)
    tipo de período    -> catalogos.PERIODOS; las predicciones NUNCA mezclan
                          períodos distintos (vacaciones ≠ semestre)

Para agregar un tipo de documento nuevo: escribir su lector (función que recibe
una ruta y devuelve un `lector.Reporte` o equivalente), registrarlo aquí con su
firma, y decidir qué vistas lo consumen. Un archivo que no coincide con ninguna
firma se rechaza con la lista de tipos soportados, en lugar de intentar leerlo
con el lector equivocado.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook

from . import lector
from .catalogos import normalizar


@dataclass(frozen=True)
class TipoDocumento:
    clave: str
    nombre: str                       # como aparece en el sistema de Centro de Cálculo
    firmas: tuple[str, ...]           # textos que identifican el membrete
    hojas: tuple[str, ...]            # hojas características (respaldo si no hay membrete)
    leer: Callable[[Path, str | None], lector.Reporte]
    alimenta_predicciones: bool


ASIGNADOS = TipoDocumento(
    clave="ASIGNADOS_CONDENSADO",
    nombre="Contadores asignados (Condensado)",
    firmas=("TOTAL DE ASIGNADOS",),
    hojas=("CLASE", "LABORATORIO"),
    leer=lector.leer_reporte,
    alimenta_predicciones=True,
)

TIPOS: dict[str, TipoDocumento] = {t.clave: t for t in (ASIGNADOS,)}


def nombres_soportados() -> str:
    return ", ".join(f"«{t.nombre}»" for t in TIPOS.values())


def identificar(ruta: Path) -> TipoDocumento | None:
    """Mira solo las primeras filas de cada hoja (rápido) y devuelve el tipo.

    Se llama DESPUÉS de lector.inspeccionar_archivo, así que aquí el archivo ya
    es un .xlsx legible. Ante cualquier problema devuelve None.
    """
    try:
        with Path(ruta).open("rb") as flujo:
            wb = load_workbook(flujo, read_only=True, data_only=True)
            try:
                hojas = {normalizar(n) for n in wb.sheetnames}
                textos = []
                for ws in wb.worksheets[:6]:
                    for fila in ws.iter_rows(min_row=1, max_row=12, max_col=5, values_only=True):
                        textos.extend(normalizar(c) for c in fila if c is not None)
            finally:
                wb.close()
    except Exception:  # noqa: BLE001
        return None
    for tipo in TIPOS.values():
        if any(f in t for f in tipo.firmas for t in textos):
            return tipo
        if set(tipo.hojas) & hojas:
            return tipo     # trae las hojas del tipo: su lector explicará qué falta
    return None
